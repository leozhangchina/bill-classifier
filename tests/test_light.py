"""回归检查只使用临时生成的示例，不包含个人账单。"""
import csv
import io
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR))
import classify_bill_light as bill


class LightClassifierTests(unittest.TestCase):
    def run_cli(self, input_path, output_path):
        return subprocess.run(
            [sys.executable, str(PROJECT_DIR / 'classify_bill_light.py'),
             str(input_path), '--output', str(output_path)],
            cwd=input_path.parent, capture_output=True, text=True, encoding='utf-8',
        )

    def read_output(self, path):
        workbook = load_workbook(path, read_only=True)
        try:
            rows = list(workbook.active.values)
            self.assertEqual(list(rows[0]), bill.OUTPUT_HEADERS)
            return rows[1:]
        finally:
            workbook.close()

    def make_wechat(self, path):
        workbook = Workbook()
        sheet = workbook.active
        for _ in range(17):
            sheet.append(['账单说明'])
        sheet.append(['交易时间', '交易类型', '交易对方', '商品', '收/支',
                      '金额(元)', '支付方式', '当前状态', '备注'])
        sheet.append(['2026-01-01 10:00:00', '转入零钱通-来自零钱', '/', '/', '/',
                      '￥100.00', '零钱', '支付成功', '/'])
        sheet.append(['2026-01-01 11:00:00', '商户消费', '测试未知商户', '无匹配商品',
                      '支出', '-12.50', '零钱', '支付成功', '/'])
        workbook.save(path)
        workbook.close()

    def test_wechat_from_other_directory(self):
        with tempfile.TemporaryDirectory() as folder:
            input_path = Path(folder) / '微信样本.xlsx'
            output_path = Path(folder) / '结果.xlsx'
            self.make_wechat(input_path)
            completed = self.run_cli(input_path, output_path)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            rows = self.read_output(output_path)
            self.assertEqual(len(rows), 2)
            self.assertEqual((rows[0][1], rows[0][2], rows[0][5], rows[0][6]),
                             ('不计收支', 100, '其他转账', '否'))
            self.assertEqual((rows[1][1], rows[1][2], rows[1][5], rows[1][6]),
                             ('支出', 12.5, '其他支出', '是'))

    def test_alipay_gb18030_yuebao_income(self):
        with tempfile.TemporaryDirectory() as folder:
            text = io.StringIO(newline='')
            writer = csv.writer(text)
            for _ in range(23):
                writer.writerow(['支付宝账单说明'])
            writer.writerow(['交易时间', '交易分类', '交易对方', '商品说明', '收/支',
                             '金额', '收/付款方式', '交易状态', '备注', ''])
            writer.writerow(['2026-01-02 10:00:00', '投资理财', '余额宝',
                             '余额宝-收益发放', '不计收支', '0.12', '余额', '成功', '', ''])
            input_path = Path(folder) / '支付宝样本.csv'
            input_path.write_bytes(text.getvalue().encode('gb18030'))
            output_path = Path(folder) / '结果.xlsx'
            completed = self.run_cli(input_path, output_path)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            row, = self.read_output(output_path)
            self.assertEqual((row[1], row[2], row[5], row[6]),
                             ('收入', 0.12, '利息收入', '否'))

    def test_override_regular_transfer_and_refund(self):
        config = bill.load_config(bill.DEFAULT_CATEGORIES)
        bill.load_overrides(bill.DEFAULT_OVERRIDES, config)
        result = bill.classify({'counterparty': '测试球馆分店', 'direction': '支出'},
                               config, {'测试球馆': '运动健身', '测试': '食品'})
        self.assertEqual(result.category, '运动健身')
        self.assertEqual(bill.transaction_flow({'type': '转入银行卡', 'direction': '/'}), 'income')
        self.assertEqual(bill.classify({'type': '退款', 'direction': '收入'}, config, {}).category, '退款')

    def test_original_bill_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as folder:
            input_path = Path(folder) / '微信样本.xlsx'
            self.make_wechat(input_path)
            original = input_path.read_bytes()
            completed = self.run_cli(input_path, input_path)
            self.assertNotEqual(completed.returncode, 0)
            self.assertIn('输出路径不能', completed.stderr)
            self.assertEqual(input_path.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
